from __future__ import annotations

from collections import Counter
import csv
import json
import os
from pathlib import Path
import tempfile
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from .models import DatasetRecord


CSV_FIELDS = (
    "sa_id",
    "year",
    "announced",
    "topic",
    "category",
    "module",
    "scope_tier",
    "is_priority",
    "cve",
    "cwe",
    "affects",
    "git_hashes",
    "svn_revisions",
    "github_commit_urls",
    "diff_available",
    "data_quality",
    "advisory_raw_url",
    "source",
    "upstream_path",
    "note",
    "commit_grade",
)

MAIN_HEADERS = (
    "SA ID",
    "Year",
    "Announced",
    "Topic",
    "Category",
    "Module",
    "Scope",
    "Priority",
    "CVE(s)",
    "Affects",
    "Git Hash",
    "SVN",
    "GitHub Commit",
    "Quality",
    "Grade",
    "Source",
    "Note",
)

PRIORITY_HEADERS = (
    "SA ID",
    "Year",
    "Topic",
    "Module",
    "Scope",
    "CVE(s)",
    "Git Hash",
    "GitHub Commit",
    "Quality",
    "Grade",
    "Source",
    "Note",
)

MAIN_WIDTHS = (22, 6, 11, 40, 10, 12, 11, 8, 26, 20, 26, 12, 50, 8, 18, 14, 28)
PRIORITY_WIDTHS = (22, 6, 44, 12, 10, 26, 26, 50, 8, 18, 12, 26)
THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)

ROW_FILLS = {
    (False, "lib"): "E8F5E9",
    (False, "tool"): "E3F2FD",
    (False, "kernel"): "FFFDE7",
    (False, "other"): "F5F5F5",
    (True, "lib"): "A5D6A7",
    (True, "tool"): "90CAF9",
}
QUALITY_FILLS = {3: "C8E6C9", 2: "FFF9C4", 1: "FFCDD2"}
GRADE_FILLS = {
    "focused": "E8F5E9",
    "medium": "FFF9C4",
    "large_concentrated": "FFE0B2",
    "bundled_multi_module": "FCE4EC",
    "noise_only": "FFCDD2",
}
SCOPE_FONTS = {"lib": "2E7D32", "tool": "1565C0", "kernel": "F57F17", "other": "455A64"}


def write_outputs(records: list[DatasetRecord], output_dir: Path) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    targets = {
        "json": output_dir / "dataset_index.json",
        "csv": output_dir / "dataset_index.csv",
        "xlsx": output_dir / "FreeBSD_SA_Dataset_Index.xlsx",
    }
    temporary: list[Path] = []
    try:
        json_temp = _temporary_path(output_dir, ".json")
        temporary.append(json_temp)
        json_temp.write_text(
            json.dumps([record.to_dict() for record in records], ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

        csv_temp = _temporary_path(output_dir, ".csv")
        temporary.append(csv_temp)
        with csv_temp.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=CSV_FIELDS, lineterminator="\r\n")
            writer.writeheader()
            for record in records:
                writer.writerow(csv_row(record))

        xlsx_temp = _temporary_path(output_dir, ".xlsx")
        temporary.append(xlsx_temp)
        workbook = build_workbook(records)
        workbook.save(xlsx_temp)

        os.replace(json_temp, targets["json"])
        os.replace(csv_temp, targets["csv"])
        os.replace(xlsx_temp, targets["xlsx"])
        return targets
    finally:
        for path in temporary:
            if path.exists():
                path.unlink()


def csv_row(record: DatasetRecord) -> dict[str, Any]:
    payload = record.to_dict()
    result: dict[str, Any] = {}
    for name in CSV_FIELDS:
        value = payload.get(name, "")
        if isinstance(value, list):
            result[name] = "; ".join(str(item) for item in value)
        elif name in {"is_priority", "diff_available"}:
            result[name] = "Yes" if value else ""
        elif value is None:
            result[name] = ""
        else:
            result[name] = value
    return result


def build_workbook(records: list[DatasetRecord]) -> Workbook:
    workbook = Workbook()
    main = workbook.active
    main.title = "SA Dataset Index"
    _write_record_sheet(
        main, records, MAIN_HEADERS, MAIN_WIDTHS, priority_layout=False, header_color="1F4E79"
    )

    priority = workbook.create_sheet("Priority lib+tool")
    _write_record_sheet(
        priority,
        [record for record in records if record.is_priority],
        PRIORITY_HEADERS,
        PRIORITY_WIDTHS,
        priority_layout=True,
        header_color="2E7D32",
    )

    kernel = workbook.create_sheet("Kernel Track")
    _write_record_sheet(
        kernel,
        [record for record in records if record.scope_tier == "kernel"],
        MAIN_HEADERS,
        MAIN_WIDTHS,
        priority_layout=False,
        header_color="6D4C41",
    )
    _write_stats_sheet(workbook.create_sheet("Stats & Legend"), records)
    return workbook


def _write_record_sheet(
    sheet: Worksheet,
    records: Iterable[DatasetRecord],
    headers: tuple[str, ...],
    widths: tuple[int, ...],
    *,
    priority_layout: bool,
    header_color: str,
) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=header_color)
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{_column_letter(len(headers))}1"
    sheet.row_dimensions[1].height = 28
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[_column_letter(index)].width = width

    for record in records:
        values = priority_values(record) if priority_layout else main_values(record)
        sheet.append(values)
        row = sheet.max_row
        _style_record_row(sheet, row, record, priority_layout=priority_layout)


def main_values(record: DatasetRecord) -> tuple[Any, ...]:
    return (
        record.sa_id,
        record.year,
        record.announced,
        record.topic,
        record.category,
        record.module,
        record.scope_tier,
        "Yes" if record.is_priority else "",
        "; ".join(record.cve),
        record.affects,
        "; ".join(record.git_hashes),
        "; ".join(record.svn_revisions),
        "; ".join(record.github_commit_urls),
        record.data_quality,
        record.commit_grade,
        record.source or "",
        record.note or "",
    )


def priority_values(record: DatasetRecord) -> tuple[Any, ...]:
    return (
        record.sa_id,
        record.year,
        record.topic,
        record.module,
        record.scope_tier,
        "; ".join(record.cve),
        "; ".join(record.git_hashes),
        "; ".join(record.github_commit_urls),
        record.data_quality,
        record.commit_grade,
        record.source or "",
        record.note or "",
    )


def _style_record_row(
    sheet: Worksheet, row: int, record: DatasetRecord, *, priority_layout: bool
) -> None:
    is_upstream = record.source == "upstream_commit"
    base_color = ROW_FILLS.get((is_upstream, record.scope_tier), "F5F5F5")
    for cell in sheet[row]:
        cell.font = Font(name="Arial", size=9)
        cell.fill = PatternFill("solid", fgColor=base_color)
        cell.border = THIN_BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    scope_column = 5 if priority_layout else 7
    quality_column = 9 if priority_layout else 14
    grade_column = 10 if priority_layout else 15
    source_column = 11 if priority_layout else 16
    note_column = 12 if priority_layout else 17
    commit_column = 8 if priority_layout else 13

    scope_cell = sheet.cell(row, scope_column)
    scope_cell.font = Font(
        name="Arial", size=9, bold=True, color=SCOPE_FONTS.get(record.scope_tier, "455A64")
    )
    if not priority_layout and record.is_priority:
        sheet.cell(row, 8).font = Font(name="Arial", size=9, bold=True, color="1B5E20")

    quality_cell = sheet.cell(row, quality_column)
    quality_cell.fill = PatternFill("solid", fgColor=QUALITY_FILLS[record.data_quality])
    quality_cell.font = Font(name="Arial", size=9, bold=True)
    quality_cell.alignment = Alignment(horizontal="center", vertical="top")

    if record.commit_grade:
        grade_cell = sheet.cell(row, grade_column)
        grade_cell.fill = PatternFill("solid", fgColor=GRADE_FILLS[record.commit_grade])
        grade_cell.font = Font(name="Arial", size=9, bold=True)
        grade_cell.alignment = Alignment(horizontal="center", vertical="top")

    if is_upstream:
        sheet.cell(row, source_column).font = Font(
            name="Arial", size=9, italic=True, color="1565C0"
        )
    if record.note:
        sheet.cell(row, note_column).font = Font(
            name="Arial", size=9, italic=True, color="6A1B9A"
        )
    if record.github_commit_urls:
        sheet.cell(row, commit_column).hyperlink = record.github_commit_urls[0]


def _write_stats_sheet(sheet: Worksheet, records: list[DatasetRecord]) -> None:
    scopes = Counter(record.scope_tier for record in records)
    qualities = Counter(record.data_quality for record in records)
    grades = Counter(record.commit_grade for record in records if record.commit_grade)
    entries = {
        1: ("数据集统计", None, None),
        2: ("总记录数", len(records), None),
        3: ("lib+tool（优先级）", sum(record.is_priority for record in records), None),
        4: ("kernel（第二阶段）", scopes["kernel"], None),
        5: ("FreeBSD-SA 来源", sum(record.source is None for record in records), None),
        6: ("Upstream 来源", sum(record.source == "upstream_commit" for record in records), None),
        7: ("Q3 全量", qualities[3], None),
        8: ("Q2 全量", qualities[2], None),
        9: ("Q3 lib+tool", sum(record.is_priority and record.data_quality == 3 for record in records), None),
        10: ("Q2 lib+tool", sum(record.is_priority and record.data_quality == 2 for record in records), None),
        13: ("commit_grade 分布（lib+tool Q3）", None, None),
        14: ("focused", grades["focused"], "直接用整体 diff"),
        15: ("medium", grades["medium"], "只取 module 目录文件"),
        16: ("large_concentrated", grades["large_concentrated"], "只取 module 目录文件（大版本升级）"),
        17: ("bundled_multi_module", grades["bundled_multi_module"], "按 module 路径过滤后提取"),
        18: ("noise_only", grades["noise_only"], "需补录正确 commit"),
        21: ("颜色图例（行背景）", None, None),
        22: ("浅绿", None, "FreeBSD-SA, lib"),
        23: ("浅蓝", None, "FreeBSD-SA, tool"),
        24: ("亮绿", None, "Upstream, lib"),
        25: ("亮蓝", None, "Upstream, tool"),
        26: ("浅黄", None, "kernel 层"),
        27: ("Q3绿", None, "Quality 3"),
        28: ("Q2黄", None, "Quality 2"),
        29: ("Q1红", None, "Quality 1"),
    }
    for row, values in entries.items():
        for column, value in enumerate(values, 1):
            if value is not None:
                sheet.cell(row, column, value)
    for column, width in enumerate((30, 14, 28), 1):
        sheet.column_dimensions[_column_letter(column)].width = width

    for row in range(1, 30):
        for cell in sheet[row]:
            cell.font = Font(name="Arial", size=10, color="000000")
            cell.alignment = Alignment(horizontal="left")
    for row in (1, 13, 21):
        sheet.cell(row, 1).font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor="1F4E79" if row == 1 else "37474F")

    legend_fills = {
        14: "E8F5E9",
        15: "FFF9C4",
        16: "FFE0B2",
        17: "FCE4EC",
        18: "FFCDD2",
        22: "E8F5E9",
        23: "E3F2FD",
        24: "A5D6A7",
        25: "90CAF9",
        26: "FFFDE7",
        27: "C8E6C9",
        28: "FFF9C4",
        29: "FFCDD2",
    }
    for row, color in legend_fills.items():
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor=color)


def _temporary_path(directory: Path, suffix: str) -> Path:
    descriptor, name = tempfile.mkstemp(prefix=".freevrg-dataset-", suffix=suffix, dir=directory)
    os.close(descriptor)
    return Path(name)


def _column_letter(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result
