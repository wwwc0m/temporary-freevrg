from __future__ import annotations

import csv
import json
from pathlib import Path
import shutil
import tempfile
from typing import Any
from zipfile import ZipFile

from openpyxl import load_workbook


REQUIRED_REFERENCE_FILES = {
    "dataset_index.json",
    "dataset_index.csv",
    "FreeBSD_SA_Dataset_Index.xlsx",
}


class ReferenceMismatch(AssertionError):
    pass


def validate_reference(output_dir: Path, reference_zip: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="freevrg-reference-") as directory:
        root = Path(directory)
        with ZipFile(reference_zip) as archive:
            located: dict[str, Path] = {}
            for member in archive.infolist():
                name = Path(member.filename).name
                if member.is_dir() or name not in REQUIRED_REFERENCE_FILES:
                    continue
                if name in located:
                    raise ReferenceMismatch(f"reference ZIP contains duplicate file: {name}")
                target = root / name
                with archive.open(member) as source, target.open("wb") as destination:
                    shutil.copyfileobj(source, destination)
                located[name] = target
        missing = REQUIRED_REFERENCE_FILES - set(located)
        if missing:
            raise ReferenceMismatch(f"reference ZIP is missing: {sorted(missing)}")

        _compare_json(output_dir / "dataset_index.json", located["dataset_index.json"])
        _compare_csv(output_dir / "dataset_index.csv", located["dataset_index.csv"])
        _compare_xlsx(
            output_dir / "FreeBSD_SA_Dataset_Index.xlsx",
            located["FreeBSD_SA_Dataset_Index.xlsx"],
        )


def _compare_json(actual_path: Path, expected_path: Path) -> None:
    actual = json.loads(actual_path.read_text(encoding="utf-8"))
    expected = json.loads(expected_path.read_text(encoding="utf-8"))
    if len(actual) != len(expected):
        raise ReferenceMismatch(f"JSON record count differs: actual={len(actual)} expected={len(expected)}")
    for index, (actual_record, expected_record) in enumerate(zip(actual, expected, strict=True)):
        if actual_record == expected_record:
            continue
        keys = list(dict.fromkeys([*expected_record, *actual_record]))
        for key in keys:
            if actual_record.get(key, _MISSING) != expected_record.get(key, _MISSING):
                raise ReferenceMismatch(
                    f"JSON first difference at record {index} ({expected_record.get('sa_id')}), "
                    f"field {key}: actual={actual_record.get(key, _MISSING)!r} "
                    f"expected={expected_record.get(key, _MISSING)!r}"
                )
        raise ReferenceMismatch(f"JSON key ordering differs at record {index}")


def _compare_csv(actual_path: Path, expected_path: Path) -> None:
    with actual_path.open(encoding="utf-8", newline="") as handle:
        actual = list(csv.reader(handle))
    with expected_path.open(encoding="utf-8", newline="") as handle:
        expected = list(csv.reader(handle))
    if not actual or not expected or actual[0] != expected[0]:
        raise ReferenceMismatch(
            f"CSV header differs: actual={actual[0] if actual else None} "
            f"expected={expected[0] if expected else None}"
        )
    if len(actual) != len(expected):
        raise ReferenceMismatch(f"CSV row count differs: actual={len(actual)} expected={len(expected)}")
    for row_index, (actual_row, expected_row) in enumerate(zip(actual, expected, strict=True), 1):
        if actual_row == expected_row:
            continue
        for column, (actual_value, expected_value) in enumerate(
            zip(actual_row, expected_row, strict=False)
        ):
            if actual_value != expected_value:
                name = actual[0][column] if column < len(actual[0]) else f"column {column}"
                raise ReferenceMismatch(
                    f"CSV first difference at row {row_index}, field {name}: "
                    f"actual={actual_value!r} expected={expected_value!r}"
                )


def _compare_xlsx(actual_path: Path, expected_path: Path) -> None:
    actual = load_workbook(actual_path)
    expected = load_workbook(expected_path)
    if actual.sheetnames != expected.sheetnames:
        raise ReferenceMismatch(
            f"XLSX sheet names differ: actual={actual.sheetnames} expected={expected.sheetnames}"
        )
    for name in actual.sheetnames:
        actual_sheet = actual[name]
        expected_sheet = expected[name]
        actual_size = (actual_sheet.max_row, actual_sheet.max_column)
        expected_size = (expected_sheet.max_row, expected_sheet.max_column)
        if actual_size != expected_size:
            raise ReferenceMismatch(
                f"XLSX {name} dimensions differ: actual={actual_size} expected={expected_size}"
            )
        for row in range(1, actual_sheet.max_row + 1):
            for column in range(1, actual_sheet.max_column + 1):
                actual_value = _normalize_cell(actual_sheet.cell(row, column).value)
                expected_value = _normalize_cell(expected_sheet.cell(row, column).value)
                if actual_value != expected_value:
                    coordinate = actual_sheet.cell(row, column).coordinate
                    raise ReferenceMismatch(
                        f"XLSX first value difference in {name}!{coordinate}: "
                        f"actual={actual_value!r} expected={expected_value!r}"
                    )
        if name != "Stats & Legend":
            if str(actual_sheet.freeze_panes) != str(expected_sheet.freeze_panes):
                raise ReferenceMismatch(
                    f"XLSX freeze pane differs in {name}: "
                    f"actual={actual_sheet.freeze_panes} expected={expected_sheet.freeze_panes}"
                )
            if actual_sheet.auto_filter.ref != expected_sheet.auto_filter.ref:
                raise ReferenceMismatch(
                    f"XLSX filter differs in {name}: actual={actual_sheet.auto_filter.ref} "
                    f"expected={expected_sheet.auto_filter.ref}"
                )
            _validate_hyperlinks(actual_sheet, 8 if name == "Priority lib+tool" else 13)
            if actual_sheet.row_dimensions[1].height != expected_sheet.row_dimensions[1].height:
                raise ReferenceMismatch(
                    f"XLSX header height differs in {name}: "
                    f"actual={actual_sheet.row_dimensions[1].height} "
                    f"expected={expected_sheet.row_dimensions[1].height}"
                )
        _compare_key_format(actual_sheet, expected_sheet)


def _validate_hyperlinks(sheet: Any, commit_column: int) -> None:
    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row, commit_column)
        if cell.value and cell.hyperlink is None:
            raise ReferenceMismatch(f"XLSX commit cell is not clickable: {sheet.title}!{cell.coordinate}")


def _compare_key_format(actual: Any, expected: Any) -> None:
    for column in range(1, actual.max_column + 1):
        letter = actual.cell(1, column).column_letter
        actual_width = actual.column_dimensions[letter].width
        expected_width = expected.column_dimensions[letter].width
        if actual_width != expected_width:
            raise ReferenceMismatch(
                f"XLSX column width differs in {actual.title}!{letter}: "
                f"actual={actual_width} expected={expected_width}"
            )
    for row in range(1, actual.max_row + 1):
        for column in range(1, actual.max_column + 1):
            actual_cell = actual.cell(row, column)
            expected_cell = expected.cell(row, column)
            if actual_cell.value is None and expected_cell.value is None:
                continue
            actual_style = _key_style(actual_cell)
            expected_style = _key_style(expected_cell)
            if actual_style != expected_style:
                raise ReferenceMismatch(
                    f"XLSX key format differs in {actual.title}!{actual_cell.coordinate}: "
                    f"actual={actual_style} expected={expected_style}"
                )


def _key_style(cell: Any) -> tuple[Any, ...]:
    color = cell.font.color
    font_color = None
    if color is not None:
        font_color = (color.type, color.rgb if color.type == "rgb" else color.indexed)
    fill = cell.fill.fgColor
    fill_color = (fill.type, fill.rgb if fill.type == "rgb" else fill.indexed)
    return (
        cell.fill.fill_type,
        fill_color,
        cell.font.bold,
        cell.font.italic,
        font_color,
        cell.alignment.horizontal,
        cell.alignment.vertical,
        cell.alignment.wrap_text,
    )


def _normalize_cell(value: Any) -> Any:
    return "" if value is None else value


_MISSING = object()
