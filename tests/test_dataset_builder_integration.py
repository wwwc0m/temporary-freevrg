from __future__ import annotations

from datetime import date
import hashlib
import json
import os
from pathlib import Path
import subprocess
from urllib.parse import urlencode

from openpyxl import load_workbook

from scripts.dataset_builder.builder import DatasetBuilder
from scripts.dataset_builder.exporters import write_outputs
from scripts.dataset_builder.models import DatasetRecord
from scripts.dataset_builder.sources import (
    GITHUB_PAGE_SIZE,
    UPSTREAM_PROJECTS,
    FreeBSDCommitMatcher,
    GitHubClient,
    GitRepository,
)


def _git(repository: Path, *arguments: str, env: dict[str, str] | None = None) -> str:
    result = subprocess.run(
        ["git", *arguments],
        cwd=repository,
        check=True,
        capture_output=True,
        text=True,
        env=env,
    )
    return result.stdout.strip()


def _init_repository(path: Path) -> None:
    path.mkdir(parents=True)
    _git(path, "init", "-q", "-b", "master")
    _git(path, "config", "user.name", "FreeVRG Test")
    _git(path, "config", "user.email", "freevrg@example.invalid")


def _commit_file(
    repository: Path,
    relative_path: str,
    content: str,
    message: str,
    committed_at: str,
) -> str:
    path = repository / relative_path
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")
    _git(repository, "add", relative_path)
    commit_env = {
        **os.environ,
        "GIT_AUTHOR_DATE": committed_at,
        "GIT_COMMITTER_DATE": committed_at,
    }
    _git(repository, "commit", "-q", "-m", message, env=commit_env)
    return _git(repository, "rev-parse", "HEAD")


def _advisory(identifier: str, module: str, cve: str, sha: str) -> tuple[str, str]:
    filename = f"FreeBSD-SA-{identifier}.{module}.asc"
    text = f"""Topic:          Fixture vulnerability in {module}
Category:       core
Module:         {module}
Announced:      2026-05-{identifier[-2:]}
Affects:        Supported FreeBSD releases.
CVE Name:       {cve}
VI.  Correction details
stable/14/ {sha}
VII. References
"""
    return filename, text


def _detail(sha: str, message: str, committed: str, paths: list[str]) -> dict[str, object]:
    return {
        "sha": sha,
        "commit": {
            "message": message,
            "committer": {"date": committed},
        },
        "files": [{"filename": path} for path in paths],
    }


class _FixtureGitHub:
    def __init__(self, details: dict[str, dict[str, object]], upstream_sha: str) -> None:
        self.details = details
        self.upstream_sha = upstream_sha

    def commit(self, owner: str, repository: str, sha: str) -> dict[str, object]:
        del owner, repository
        return self.details[sha]

    def commits_for_path(
        self,
        owner: str,
        repository: str,
        path: str,
        since: date,
        until: date,
    ) -> list[dict[str, str]]:
        del owner, repository, path, since, until
        return [{"sha": self.upstream_sha}]

    def search_commits(self, query: str) -> list[dict[str, str]]:
        del query
        return []


def test_hermetic_raw_sources_build_through_all_outputs(tmp_path: Path) -> None:
    cache = tmp_path / "cache"
    doc_repository = cache / "git" / "freebsd__freebsd-doc"
    _init_repository(doc_repository)
    full_shas = [character * 40 for character in "abcd"]
    advisories = [
        _advisory("26:01", "openssl", "CVE-2026-8001", full_shas[0]),
        _advisory("26:02", "kernelish", "CVE-2026-8002", full_shas[1]),
        _advisory("26:03", "openssl", "CVE-2026-8003", full_shas[2]),
        _advisory("26:04", "openssl", "CVE-2026-8004", full_shas[3]),
    ]
    for filename, text in advisories:
        target = f"website/static/security/advisories/{filename}"
        _commit_file(doc_repository, target, text, f"Add {filename}", "2026-05-20T12:00:00Z")

    for project in UPSTREAM_PROJECTS:
        repository = cache / "git" / f"{project.owner}__{project.repository}"
        _init_repository(repository)
        message = (
            "Fix CVE-2026-9001 in OpenSSL 3.2.1"
            if project.key == "openssl"
            else f"Routine {project.key} maintenance"
        )
        _commit_file(repository, "source.txt", project.key, message, "2026-04-01T12:00:00Z")

    curation_path = tmp_path / "curation.json"
    curation_path.write_text(
        json.dumps(
            {
                "version": "1.4",
                "snapshot_date": "2026-01-01",
                "scope_overrides": {
                    "FreeBSD-SA-26:03": {"value": "tool", "reason": "Fixture override."}
                },
                "commit_grade_audit": {
                    "FreeBSD-SA-26:03": {"value": "medium", "reason": "Fixture override."}
                },
            }
        ),
        encoding="utf-8",
    )

    upstream_sha = "e" * 40
    details = {
        full_shas[0][:12]: _detail(
            full_shas[0], "Fix CVE-2026-8001", "2026-05-01T00:00:00Z", ["sys/net/ssl.c"]
        ),
        full_shas[1][:12]: _detail(
            full_shas[1], "Fix CVE-2026-8002", "2026-05-02T00:00:00Z", ["lib/libc/io.c"]
        ),
        full_shas[2][:12]: _detail(
            full_shas[2], "Fix CVE-2026-8003", "2026-05-03T00:00:00Z", ["sys/kern/fix.c"]
        ),
        full_shas[3][:12]: _detail(
            full_shas[3],
            "Document CVE-2026-8004",
            "2026-05-04T00:00:00Z",
            ["release/doc/en/security.xml"],
        ),
        upstream_sha: _detail(
            upstream_sha,
            "Import OpenSSL 3.2.1 security fix for CVE-2026-9001",
            "2026-04-03T00:00:00Z",
            ["crypto/openssl/ssl/tls.c"],
        ),
    }
    builder = DatasetBuilder(
        cache_dir=cache,
        curation_path=curation_path,
        snapshot_date=date(2026, 6, 1),
        offline=True,
    )
    builder.matcher = FreeBSDCommitMatcher(_FixtureGitHub(details, upstream_sha))  # type: ignore[arg-type]

    records = builder.build()
    by_id = {record.sa_id: record for record in records}
    assert by_id["FreeBSD-SA-26:01"].scope_tier == "kernel"
    assert by_id["FreeBSD-SA-26:01"].commit_grade == "focused"
    assert by_id["FreeBSD-SA-26:02"].scope_tier == "lib"
    assert by_id["FreeBSD-SA-26:03"].scope_tier == "tool"
    assert by_id["FreeBSD-SA-26:03"].commit_grade == "medium"
    assert by_id["FreeBSD-SA-26:04"].git_hashes == []
    assert by_id["FreeBSD-SA-26:04"].commit_grade == "noise_only"
    assert by_id["FreeBSD-SA-26:04"].data_quality == 2
    upstream = by_id["UPSTREAM-OPENSSL-CVE-2026-9001"]
    assert upstream.git_hashes == [upstream_sha[:12]]
    assert upstream.github_commit_urls == [
        f"https://github.com/freebsd/freebsd-src/commit/{upstream_sha}"
    ]
    assert upstream.scope_tier == "lib"
    assert upstream.commit_grade == "focused"

    outputs = write_outputs(records, tmp_path / "output")
    json_rows = json.loads(outputs["json"].read_text(encoding="utf-8"))
    assert len(json_rows) == len(records) == 5
    assert "changed_paths" not in json_rows[0]
    assert len(outputs["csv"].read_text(encoding="utf-8").splitlines()) == 6
    workbook = load_workbook(outputs["xlsx"], read_only=True)
    assert workbook["SA Dataset Index"].max_row == 6


def _upstream_record() -> DatasetRecord:
    record = DatasetRecord(
        sa_id="UPSTREAM-OPENSSL-CVE-2026-9999",
        year=2026,
        announced="2026-01-01",
        topic="OpenSSL 3.2.1 security update",
        category="contrib",
        module="openssl",
        scope_tier="lib",
        is_priority=True,
        cve=["CVE-2026-9999"],
        source="upstream_commit",
        upstream_path="crypto/openssl",
        upstream_message="Fix CVE-2026-9999 in OpenSSL 3.2.1",
    )
    record.recompute()
    return record


class _OneCandidateClient:
    def __init__(self, detail: dict[str, object]) -> None:
        self.detail = detail

    def commits_for_path(self, *args: object) -> list[dict[str, str]]:
        del args
        return [{"sha": str(self.detail["sha"])}]

    def commit(self, *args: object) -> dict[str, object]:
        del args
        return self.detail


def test_unrelated_commit_touching_openssl_path_is_not_matched() -> None:
    sha = "f" * 40
    client = _OneCandidateClient(
        _detail(sha, "Refactor provider initialization", "2026-02-01T00:00:00Z", ["crypto/openssl/a.c"])
    )
    matcher = FreeBSDCommitMatcher(client)  # type: ignore[arg-type]
    assert matcher.match_upstream(_upstream_record(), UPSTREAM_PROJECTS[0]) is None


class _PagedClient(GitHubClient):
    def __init__(self, cache_dir: Path) -> None:
        super().__init__(cache_dir, token=None, offline=False, refresh=False)
        self.requested: list[tuple[str, int]] = []
        self.detail_requests: list[str] = []

    def get_json(
        self, url: str, parameters: dict[str, str | int] | None = None
    ) -> object:
        if parameters is None:
            sha = url.rsplit("/", 1)[-1]
            self.detail_requests.append(sha)
            message = (
                "Import OpenSSL security fix for CVE-2026-9999"
                if sha == "f" * 40
                else "Routine OpenSSL refactor"
            )
            return _detail(sha, message, "2026-02-01T00:00:00Z", ["crypto/openssl/a.c"])
        page = int(parameters["page"])
        self.requested.append((url, page))
        if page == 1:
            items = [{"sha": f"{index:040x}"} for index in range(GITHUB_PAGE_SIZE)]
        else:
            items = [{"sha": "f" * 40}]
        return {"items": items} if url.endswith("search/commits") else items


def test_commit_search_and_path_commits_reach_second_page(tmp_path: Path) -> None:
    client = _PagedClient(tmp_path)
    search = client.search_commits("repo:freebsd/freebsd-src CVE-2026-9999")
    path = client.commits_for_path(
        "freebsd", "freebsd-src", "crypto/openssl", date(2026, 1, 1), date(2026, 2, 1)
    )
    assert search[-1]["sha"] == "f" * 40
    assert path[-1]["sha"] == "f" * 40
    assert [page for url, page in client.requested if url.endswith("search/commits")] == [1, 2]
    assert [page for url, page in client.requested if not url.endswith("search/commits")] == [1, 2]

    evidence = FreeBSDCommitMatcher(client).match_upstream(_upstream_record(), UPSTREAM_PROJECTS[0])
    assert evidence is not None
    assert evidence.full_shas == ("f" * 40,)
    assert len(client.detail_requests) == GITHUB_PAGE_SIZE + 1


def test_offline_pagination_reads_every_cached_page(tmp_path: Path) -> None:
    query = "repo:freebsd/freebsd-src cached"
    base_url = "https://api.github.com/search/commits"
    for page, items in (
        (1, [{"sha": f"{index:040x}"} for index in range(GITHUB_PAGE_SIZE)]),
        (2, [{"sha": "f" * 40}]),
    ):
        parameters = {"q": query, "per_page": GITHUB_PAGE_SIZE, "page": page}
        request_url = f"{base_url}?{urlencode(parameters)}"
        cache_path = tmp_path / f"{hashlib.sha256(request_url.encode()).hexdigest()}.json"
        cache_path.write_text(json.dumps({"items": items}), encoding="utf-8")
    client = GitHubClient(tmp_path, token=None, offline=True, refresh=False)
    assert len(client.search_commits(query)) == GITHUB_PAGE_SIZE + 1


def test_online_git_cache_updates_and_deepens_for_older_snapshot(tmp_path: Path) -> None:
    remote = tmp_path / "remote"
    _init_repository(remote)
    old_sha = _commit_file(remote, "history.txt", "old", "Old", "2020-01-01T00:00:00Z")
    _commit_file(remote, "history.txt", "new", "New", "2026-01-01T00:00:00Z")
    cache = tmp_path / "git-cache"
    repository = GitRepository(
        cache,
        remote.as_uri(),
        "master",
        offline=False,
        refresh=False,
    )
    repository.ensure(date(2025, 1, 1))
    assert repository.revision_at(date(2020, 12, 31)) == old_sha

    latest_sha = _commit_file(remote, "latest.txt", "latest", "Latest", "2026-02-01T00:00:00Z")
    repository.ensure(date(2025, 1, 1))
    assert repository.revision_at(date(2026, 12, 31)) == latest_sha
    metadata = json.loads((cache / ".freevrg-cache.json").read_text(encoding="utf-8"))
    assert metadata["repository_url"] == remote.as_uri()
    assert metadata["branch"] == "master"
    assert metadata["latest_remote_commit"] == latest_sha
    assert metadata["earliest_available_commit_date"].startswith("2020-01-01")
